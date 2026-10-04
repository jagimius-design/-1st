"""Expense-Tracker.xlsx: an expense log, a category list with budgets, and a yearly summary."""

from datetime import date

from openpyxl.chart import BarChart, Reference
from openpyxl.utils import get_column_letter

from style import (BLUE, DATE, INPUT, INPUT_NOTE, MONEY, MUTED, NAVY, PALE, TOPLINE, cell,
                   dropdown, header, negative_red, sheet, table, tile, title, workbook)

YEAR = 2026
CATEGORIES = [  # name, monthly budget
    ("Software & subscriptions", 120), ("Equipment", 150), ("Office & coworking", 250),
    ("Travel", 150), ("Meals & entertainment", 80), ("Marketing & advertising", 100),
    ("Professional services", 100), ("Phone & internet", 90), ("Education & books", 50),
    ("Bank & payment fees", 30), ("Insurance", 60), ("Other", 50),
]
METHODS = ["Card", "Bank transfer", "Cash", "PayPal", "Other"]
CAT_FIRST, CAT_LAST = 5, 34  # category slots on Categories
EXP_FIRST, EXP_LAST = 5, 1004  # expense rows on Expenses
SUM_FIRST = 9  # first category row on Summary


def _sample():
    """(date, description, category, amount, method, deductible), Jan-Sep of YEAR."""
    rows = []
    for m in range(1, 10):
        rows += [
            (date(YEAR, m, 2), "Design software subscription", "Software & subscriptions", 59.99,
             "Card", "Yes"),
            (date(YEAR, m, 3), "Cloud storage", "Software & subscriptions", 11.99, "Card", "Yes"),
            (date(YEAR, m, 5), "Coworking membership", "Office & coworking", 220, "Bank transfer",
             "Yes"),
            (date(YEAR, m, 12), "Mobile & home internet", "Phone & internet", 85, "Card", "Yes"),
            (date(YEAR, m, 28), "Payment processor fees", "Bank & payment fees", 18.4 + m, "PayPal",
             "Yes"),
        ]
    rows += [
        (date(YEAR, 1, 15), "Laptop stand and keyboard", "Equipment", 189, "Card", "Yes"),
        (date(YEAR, 2, 9), "Client lunch", "Meals & entertainment", 64.5, "Card", "Yes"),
        (date(YEAR, 3, 20), "Train to client workshop", "Travel", 142, "Card", "Yes"),
        (date(YEAR, 3, 21), "Hotel, one night", "Travel", 118, "Card", "Yes"),
        (date(YEAR, 4, 4), "Accountant - annual return", "Professional services", 450,
         "Bank transfer", "Yes"),
        (date(YEAR, 4, 18), "Online course: pricing", "Education & books", 79, "Card", "Yes"),
        (date(YEAR, 5, 6), "Portfolio ads", "Marketing & advertising", 120, "Card", "Yes"),
        (date(YEAR, 5, 30), "Professional liability insurance", "Insurance", 540,
         "Bank transfer", "Yes"),
        (date(YEAR, 6, 11), "External monitor", "Equipment", 329, "Card", "Yes"),
        (date(YEAR, 7, 2), "Business cards", "Marketing & advertising", 45, "Card", "Yes"),
        (date(YEAR, 7, 22), "Coffee with prospect", "Meals & entertainment", 12.8, "Cash", "No"),
        (date(YEAR, 8, 14), "Conference ticket", "Education & books", 299, "Card", "Yes"),
        (date(YEAR, 8, 15), "Flight to conference", "Travel", 236, "Card", "Yes"),
        (date(YEAR, 9, 8), "Books", "Education & books", 42.9, "Card", "Yes"),
        (date(YEAR, 9, 19), "Desk lamp (home use)", "Other", 38, "Card", "No"),
    ]
    return sorted(rows, key=lambda r: r[0])


SAMPLE = _sample()


def _categories(wb):
    ws = sheet(wb, "Categories", [3, 32, 18], tab=BLUE)
    title(ws, "Categories", "Rename, add or remove categories here; the dropdowns and the"
          " summary follow. " + INPUT_NOTE)
    header(ws, CAT_FIRST - 1, 2, ["Category", "Monthly budget"])
    for i, r in enumerate(range(CAT_FIRST, CAT_LAST + 1)):
        name, budget = CATEGORIES[i] if i < len(CATEGORIES) else (None, None)
        ws[f"B{r}"], ws[f"C{r}"] = name, budget
    table(ws, CAT_FIRST, CAT_LAST, {"B": (None, True, "left"), "C": (MONEY, True, "right")})
    ws.freeze_panes = f"A{CAT_FIRST}"


def _expenses(wb):
    ws = sheet(wb, "Expenses", [3, 13, 36, 26, 13, 15, 12, 10, 30], tab=BLUE)
    title(ws, "Expenses", "One row per expense. " + INPUT_NOTE)
    header(ws, EXP_FIRST - 1, 2, ["Date", "Description", "Category", "Amount", "Payment method",
                                  "Tax deductible", "Receipt", "Notes"])
    for r, row in zip(range(EXP_FIRST, EXP_LAST + 1), SAMPLE):
        for col, value in zip("BCDEFG", row):
            ws[f"{col}{r}"] = value
        ws[f"H{r}"] = "Yes"
    table(ws, EXP_FIRST, EXP_LAST, {
        "B": (DATE, True, "center"), "C": (None, True, "left"), "D": (None, True, "left"),
        "E": (MONEY, True, "right"), "F": (None, True, "center"), "G": (None, True, "center"),
        "H": (None, True, "center"), "I": (None, True, "left")})
    for col, source in [("D", f"Categories!$B${CAT_FIRST}:$B${CAT_LAST}"),
                        ("F", '"' + ",".join(METHODS) + '"'), ("G", '"Yes,No"'), ("H", '"Yes,No"')]:
        dropdown(ws, f"{col}{EXP_FIRST}:{col}{EXP_LAST}", source)
    ws.freeze_panes = f"C{EXP_FIRST}"
    ws.auto_filter.ref = f"B{EXP_FIRST - 1}:I{EXP_LAST}"


def _summary(wb):
    n = CAT_LAST - CAT_FIRST + 1
    last = SUM_FIRST + n - 1
    ws = sheet(wb, "Summary", [3, 28] + [10.5] * 12 + [13, 13, 13], tab=NAVY)
    title(ws, "Expense summary", "Totals by category and month for the year below. "
          + INPUT_NOTE)
    cell(ws, "B4", "Year", bold=True, color=NAVY, align="right")
    cell(ws, "C4", YEAR, "0", bold=True, bg=INPUT, align="center")

    total_row, uncat_row, deduct_row = last + 1, last + 2, last + 3
    year_total = f"$O${total_row}"
    tile(ws, 5, 2, "SPENT THIS YEAR", f"={year_total}")
    tile(ws, 5, 4, "TAX-DEDUCTIBLE", f"=$O${deduct_row}", span=3)
    tile(ws, 5, 8, "AVERAGE PER MONTH",
         f"={year_total}/IF($C$4=YEAR(TODAY()),MONTH(TODAY()),12)", span=3)
    tile(ws, 5, 12, "CATEGORIES OVER BUDGET", f'=COUNTIF($Q${SUM_FIRST}:$Q${last},"<0")',
         fmt="0", span=3)

    months = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
    header(ws, SUM_FIRST - 1, 2, ["Category"] + months + ["Year total", "Year budget",
                                                         "Left in budget"])
    d, a, c = (f"Expenses!$B${EXP_FIRST}:$B${EXP_LAST}", f"Expenses!$E${EXP_FIRST}:$E${EXP_LAST}",
               f"Expenses!$D${EXP_FIRST}:$D${EXP_LAST}")

    def window(m):
        return f'{d},">="&DATE($C$4,{m},1),{d},"<"&DATE($C$4,{m + 1},1)'

    for i, r in enumerate(range(SUM_FIRST, last + 1)):
        src = CAT_FIRST + i
        ws[f"B{r}"] = f'=IF(Categories!$B${src}="","",Categories!$B${src})'
        for m in range(1, 13):
            col = get_column_letter(2 + m)
            ws[f"{col}{r}"] = f'=IF($B{r}="","",SUMIFS({a},{c},$B{r},{window(m)}))'
        ws[f"O{r}"] = f'=IF($B{r}="","",SUM(C{r}:N{r}))'
        ws[f"P{r}"] = f'=IF(OR($B{r}="",Categories!$C${src}=""),"",Categories!$C${src}*12)'
        ws[f"Q{r}"] = f'=IF(P{r}="","",P{r}-O{r})'
    body = {"B": (None, False, "left")}
    body.update({get_column_letter(k): (MONEY, False, "right") for k in range(3, 18)})
    table(ws, SUM_FIRST, last, body)
    negative_red(ws, f"Q{SUM_FIRST}:Q{last}", f"Q{SUM_FIRST}")

    labels = {total_row: "Total", uncat_row: "Not in a category", deduct_row: "Tax-deductible"}
    for r, label in labels.items():
        cell(ws, f"B{r}", label, bold=r == total_row, color=NAVY, bg=PALE)
        for k in range(3, 18):
            col = get_column_letter(k)
            if k <= 14:
                m = k - 2
                f = {total_row: f"=SUMIFS({a},{window(m)})",
                     uncat_row: f"={col}{total_row}-SUM({col}{SUM_FIRST}:{col}{last})",
                     deduct_row: f'=SUMIFS({a},Expenses!$G${EXP_FIRST}:$G${EXP_LAST},"Yes",'
                                 f'{window(m)})'}[r]
            elif k == 15:
                f = f"=SUM(C{r}:N{r})"
            elif k == 16 and r == total_row:
                f = f"=SUM(P{SUM_FIRST}:P{last})"
            elif k == 17 and r == total_row:
                f = f"=P{r}-O{r}"
            else:
                f = None
            cell(ws, f"{col}{r}", f, MONEY, bold=r == total_row, color=NAVY, bg=PALE,
                 align="right", border=TOPLINE if r == total_row else None)
    ws[f"B{total_row}"].border = TOPLINE
    cell(ws, f"B{deduct_row + 1}", '"Not in a category" catches expenses whose category is'
         " blank or misspelled; it should read zero.", size=9, color=MUTED, italic=True)
    ws.freeze_panes = f"C{SUM_FIRST}"

    chart = BarChart()
    chart.title = "Spending by month"
    chart.style = 10
    chart.legend = None
    chart.y_axis.numFmt = "#,##0"
    chart.y_axis.majorGridlines = None
    chart.add_data(Reference(ws, min_col=3, max_col=14, min_row=total_row), from_rows=True)
    chart.set_categories(Reference(ws, min_col=3, max_col=14, min_row=SUM_FIRST - 1))
    chart.series[0].graphicalProperties.solidFill = "3E7CB1"
    chart.series[0].graphicalProperties.line.solidFill = "3E7CB1"
    chart.height, chart.width = 7.5, 24
    ws.add_chart(chart, f"B{deduct_row + 3}")
    return ws


def build(path):
    wb = workbook()
    _summary(wb)
    _expenses(wb)
    _categories(wb)
    wb.save(path)
