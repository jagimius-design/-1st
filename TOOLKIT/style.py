"""Shared look of the templates: palette, cell styles and layout helpers."""

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.properties import PageSetupProperties

INPUT_NOTE = "Yellow cells are yours to fill in. Everything else calculates."

NAVY = "1F3A5F"
BLUE = "3E7CB1"
PALE = "EEF3F8"
INPUT = "FFF8DC"
LINE = "D5DCE4"
MUTED = "6B7785"
INK = "1F2933"
RED_BG, RED_FG = "FDE2E1", "9B1C1C"
GREEN_BG, GREEN_FG = "E3F4E8", "1E6B3A"
AMBER_BG, AMBER_FG = "FFF1D6", "8A5A00"

FONT = "Calibri"
MONEY = '"£"#,##0.00;-"£"#,##0.00;"-"'
PCT = "0.0%"
DATE = "dd/mm/yyyy"
HOURS = "0.00"

_thin = Side(style="thin", color=LINE)
GRID = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)
UNDERLINE = Border(bottom=Side(style="medium", color=NAVY))
TOPLINE = Border(top=Side(style="thin", color=NAVY))


def font(size=10, bold=False, color=INK, italic=False):
    return Font(name=FONT, size=size, bold=bold, color=color or INK, italic=italic)


def fill(color):
    return PatternFill("solid", start_color=color, end_color=color)


def workbook():
    wb = Workbook()
    wb.remove(wb.active)
    wb.calculation.fullCalcOnLoad = True  # openpyxl stores no results
    return wb


def sheet(wb, name, widths, tab=NAVY, gridlines=False):
    """A sheet with column widths, A first."""
    ws = wb.create_sheet(name)
    ws.sheet_properties.tabColor = tab
    ws.sheet_view.showGridLines = gridlines
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    return ws


def title(ws, text, subtitle, col=2):
    ws.row_dimensions[2].height = 30
    c = ws.cell(2, col, text)
    c.font = font(20, True, NAVY)
    c.alignment = Alignment(vertical="center")
    s = ws.cell(3, col, subtitle)
    s.font = font(10, color=MUTED, italic=True)


def header(ws, row, col, labels):
    ws.row_dimensions[row].height = 30
    for i, label in enumerate(labels):
        c = ws.cell(row, col + i, label)
        c.font = font(10, True, "FFFFFF")
        c.fill = fill(NAVY)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = GRID


def cell(ws, ref, value=None, fmt=None, bold=False, size=10, color=INK, bg=None,
         align=None, border=None, italic=False):
    c = ws[ref]
    if value is not None:
        c.value = value
    c.font = font(size, bold, color, italic)
    if fmt:
        c.number_format = fmt
    if bg:
        c.fill = fill(bg)
    if align:
        c.alignment = Alignment(horizontal=align, vertical="center")
    if border:
        c.border = border
    return c


def table(ws, first_row, last_row, columns):
    """Style a table body. `columns` maps column letter -> (number format, is_input, align)."""
    for r in range(first_row, last_row + 1):
        band = PALE if (r - first_row) % 2 else None
        for col, (fmt, is_input, align) in columns.items():
            c = ws[f"{col}{r}"]
            c.font = font()
            c.border = GRID
            c.alignment = Alignment(horizontal=align, vertical="center")
            if fmt:
                c.number_format = fmt
            if is_input:
                c.fill = fill(INPUT)
            elif band:
                c.fill = fill(band)


def tile(ws, row, col, label, formula, fmt=MONEY, span=1):
    """A KPI tile `span` columns wide: small label over a large figure."""
    for k in range(col, col + span):
        for r in (row, row + 1):
            ws.cell(r, k).fill = fill(PALE)
        ws.cell(row + 1, k).border = Border(bottom=Side(style="thick", color=BLUE))
    if span > 1:
        ws.merge_cells(start_row=row, start_column=col, end_row=row, end_column=col + span - 1)
        ws.merge_cells(start_row=row + 1, start_column=col, end_row=row + 1,
                       end_column=col + span - 1)
    top = ws.cell(row, col, label)
    top.font = font(9, True, MUTED)
    top.alignment = Alignment(horizontal="left", vertical="bottom", indent=1)
    big = ws.cell(row + 1, col, formula)
    big.font = font(18, True, NAVY)
    big.number_format = fmt
    big.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[row].height = 20
    ws.row_dimensions[row + 1].height = 34


def dropdown(ws, rng, source):
    """List validation; `source` is a sheet range or a quoted comma list."""
    dv = DataValidation(type="list", formula1=source, allow_blank=True, showDropDown=False)
    dv.error = "Pick a value from the list."
    dv.errorTitle = "Not in the list"
    ws.add_data_validation(dv)
    dv.add(rng)


def highlight(ws, rng, condition, bg, fg):
    """Conditional fill for `rng`; `condition` is written for its top-left cell."""
    ws.conditional_formatting.add(rng, FormulaRule(
        formula=[condition], fill=fill(bg), font=Font(color=fg, bold=True)))


def status_colors(ws, rng, first_cell, states):
    """`states` maps a cell's text to (background, foreground)."""
    for text, (bg, fg) in states.items():
        highlight(ws, rng, f'{first_cell}="{text}"', bg, fg)


def negative_red(ws, rng, first_cell):
    highlight(ws, rng, f"AND(ISNUMBER({first_cell}),{first_cell}<0)", RED_BG, RED_FG)
